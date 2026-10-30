from contextlib import asynccontextmanager
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Literal
import shutil, zipfile, io, json, threading, ipaddress, time as time_module, hmac, os
import unicodedata, re, uuid
if os.name=='nt': import msvcrt as _file_lock
else: import fcntl as _file_lock
from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Query, Header, Request, Response
from fastapi.responses import HTMLResponse, StreamingResponse, FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from starlette.concurrency import run_in_threadpool
import logging
from sqlalchemy.orm import Session
from sqlalchemy import or_, func, inspect
from pydantic import BaseModel, ConfigDict, Field, SecretStr, field_validator
from openpyxl import load_workbook, Workbook
from .database import Base, engine, get_db, DATA_DIR, BACKUP_DIR, IMPORT_BACKUP_DIR, DB_PATH, JSON_DB_PATH, SessionLocal, DATA_LOCK, _backup_path_notice, initialize_json_database, acquire_data_lock, release_data_lock, DatabaseBusyError, release_data_lock, DatabaseBusyError, SLOW_LOCK_SECONDS
from .models import Product, Movement, Customer, Supplier, BankAccount, AccountBalanceAdjustment, AppUser, BankTransaction, ServiceOrder, Order, OrderItem, PurchaseReceipt, ServiceOrderPhoto, Guarantee, GuaranteeOccurrence, FinancialEntry, ImportLog, Sale, SaleItem, ServicePart, ServiceOrderEvent
from .services import backup_service
from . import security, authorization, guarantees as guarantee_rules
from .validation import normalize_phone, normalize_cpf, normalize_serial
from . import customers as customer_rules

_instance_lock_handle=None

_lmf_logger=logging.getLogger('lmf')
if not _lmf_logger.handlers:
    _handler=logging.StreamHandler()
    _handler.setFormatter(logging.Formatter('%(asctime)s %(levelname)s %(name)s: %(message)s'))
    _lmf_logger.addHandler(_handler)
    _lmf_logger.setLevel(os.getenv('LMF_LOG_LEVEL','INFO').upper())
    _lmf_logger.propagate=False
api_logger=logging.getLogger('lmf.api')

def _load_auth_user(user_id:int):
    """Read the signed-in user in a worker thread so the event loop is never blocked by DATA_LOCK."""
    acquired_at=acquire_data_lock('auth-middleware')
    try:
        with SessionLocal() as db:
            user=db.get(AppUser,user_id)
            return {'id':user.id,'username':user.username,'is_admin':user.is_admin,'active':user.active} if user else None
    finally:
        release_data_lock('auth-middleware',acquired_at)

def _audit_outcome(request:Request,status_code:int,response=None):
    """Se a requisição consumiu uma autorização administrativa, registra se a operação foi executada ou falhou.
    Roda DEPOIS de a resposta sair: o bloqueio do banco só é liberado após o envio, então esperar por ele
    antes de responder travaria a requisição (mesmo padrão do deadlock original)."""
    operation=getattr(request.state,'audit_operation',None)
    if not operation: return
    args=(operation[0],operation[1],getattr(request.state,'user_id',None),getattr(request.state,'username',None),authorization.client_ip(request),status_code)
    authorization.start_outcome_thread(*args)

def _acquire_instance_lock():
    """Prevent multiple server processes from writing separate caches to one JSON file."""
    global _instance_lock_handle
    path=DATA_DIR/'lmf_control.instance.lock'
    path.parent.mkdir(parents=True,exist_ok=True)
    handle=path.open('a+b')
    handle.seek(0,os.SEEK_END)
    if handle.tell()==0:
        handle.write(b'0'); handle.flush()
    handle.seek(0)
    try:
        if os.name=='nt': _file_lock.locking(handle.fileno(),_file_lock.LK_NBLCK,1)
        else: _file_lock.flock(handle.fileno(),_file_lock.LOCK_EX|_file_lock.LOCK_NB)
    except OSError as error:
        handle.close()
        raise RuntimeError('Outra instância do LMF Control já está usando este banco JSON. Feche o outro servidor antes de iniciar.') from error
    _instance_lock_handle=handle

def _release_instance_lock():
    global _instance_lock_handle
    if _instance_lock_handle is None: return
    _instance_lock_handle.seek(0)
    try:
        if os.name=='nt': _file_lock.locking(_instance_lock_handle.fileno(),_file_lock.LK_UNLCK,1)
        else: _file_lock.flock(_instance_lock_handle.fileno(),_file_lock.LOCK_UN)
    finally:
        _instance_lock_handle.close()
        _instance_lock_handle=None

@asynccontextmanager
async def lifespan(app):
    _acquire_instance_lock()
    initialize_json_database()
    with SessionLocal() as db:
        if not db.query(BankAccount).filter(BankAccount.business_unit=='Assistência',BankAccount.provider=='InfinitePay').first():
            db.add(BankAccount(name='Conta da Assistência',business_unit='Assistência',provider='InfinitePay',opening_balance=0,notes='Conta de controle local. Preencha o saldo inicial real antes de usar o saldo calculado.'))
            db.commit()
        customer_rules.assign_missing_public_seq(db)
        admin=db.query(AppUser).filter(AppUser.is_admin==True,AppUser.active==True).order_by(AppUser.id).first()
        if admin:
            try: recovery_code=security.RECOVERY_CODE_PATH.read_text(encoding='utf-8').strip()
            except OSError: recovery_code=''
            if not admin.recovery_code_hash or not security.verify_password(recovery_code,admin.recovery_code_hash):
                recovery_code=security.new_recovery_code()
                admin.recovery_code_hash=security.password_hash(recovery_code)
                db.commit()
                security.save_recovery_code(recovery_code)
    stop=threading.Event()
    def daily_backup():
        while not stop.is_set():
            settings=backup_service.read_settings()
            try: hour,minute=map(int,settings.get('backup_time','02:00').split(':'))
            except Exception: hour,minute=2,0
            now=datetime.now(); next_run=datetime.combine(now.date(),time(hour,minute))
            if next_run<=now: next_run+=timedelta(days=1)
            if stop.wait(min(30,max(1,(next_run-now).total_seconds()))): break
            if datetime.now()>=next_run:
                try: backup_service.create_backup('automatico')
                except Exception as error: backup_service.set_status('falhou',str(error))
    thread=threading.Thread(target=daily_backup,name='lmf-daily-backup',daemon=True); thread.start()
    def guarantee_expiry():
        # Expira garantias vencidas sem ninguém abrir a tela: ao iniciar, depois a cada 5 min e logo após a meia-noite.
        while True:
            acquired=None
            try:
                acquired=acquire_data_lock('expirar-garantias',timeout=10)
                with SessionLocal() as session: guarantee_rules.expire_due_guarantees(session)
            except DatabaseBusyError: pass
            except Exception: logging.getLogger('lmf').exception('Falha ao expirar garantias')
            finally:
                if acquired is not None: release_data_lock('expirar-garantias',acquired)
            now=datetime.now(); midnight=datetime.combine(now.date()+timedelta(days=1),time(0,0,5))
            if stop.wait(min(300,max(1,(midnight-now).total_seconds()))): break
    expiry_thread=threading.Thread(target=guarantee_expiry,name='lmf-guarantee-expiry',daemon=True); expiry_thread.start()
    yield
    stop.set(); thread.join(timeout=2); expiry_thread.join(timeout=12)
    authorization.drain_pending(10)  # não perder o registro do resultado de uma operação em andamento
    _release_instance_lock()

app = FastAPI(title='LMF Control', version='1.0.0', lifespan=lifespan)
app.mount('/static', StaticFiles(directory=Path(__file__).parent / 'static'), name='static')

AUTH_PUBLIC_PATHS={'/api/auth/status','/api/auth/setup','/api/auth/login','/api/auth/forgot-password'}
_login_attempts:dict[str,list[float]]={}
_recovery_attempts:dict[str,list[float]]={}

@app.middleware('http')
async def require_localhost_and_auth(request:Request,call_next):
    if os.getenv('LMF_ALLOW_REMOTE','').lower() not in ('1','true','yes'):
        remote=request.client.host if request.client else ''
        try: loopback=ipaddress.ip_address(remote).is_loopback
        except ValueError: loopback=remote=='testclient'
        if not loopback: return JSONResponse({'detail':'Acesso permitido somente por este computador.'},status_code=403)
    if request.url.path.startswith('/api/') and request.url.path not in AUTH_PUBLIC_PATHS:
        session=security.read_session(request.cookies.get('lmf_session'))
        if not session: return JSONResponse({'detail':'Autenticação necessária.'},status_code=401)
        try: user=await run_in_threadpool(_load_auth_user,session['uid'])
        except DatabaseBusyError as error:
            return JSONResponse({'detail':f'{error} Tente novamente em instantes.'},status_code=503)
        if not user or not user['active']: return JSONResponse({'detail':'Sessão inválida. Entre novamente.'},status_code=401)
        if request.method not in ('GET','HEAD','OPTIONS') and not hmac.compare_digest(request.headers.get('X-CSRF-Token',''),session['csrf']):
            return JSONResponse({'detail':'Token CSRF ausente ou inválido.'},status_code=403)
        request.state.user_id=user['id']
        request.state.username=user['username']
        request.state.is_admin=user['is_admin']
    started=time_module.monotonic()
    try:
        response=await call_next(request)
    except Exception:
        api_logger.exception('Falha ao processar %s %s após %.2fs',request.method,request.url.path,time_module.monotonic()-started)
        _audit_outcome(request,500)
        raise
    _audit_outcome(request,response.status_code,response)
    elapsed=time_module.monotonic()-started
    if elapsed>=SLOW_LOCK_SECONDS:
        api_logger.warning('Requisição lenta: %s %s respondeu %s em %.2fs',request.method,request.url.path,response.status_code,elapsed)
    return response

class AuthIn(BaseModel):
    username:str=Field(min_length=1,max_length=60)
    password:str=Field(min_length=6,max_length=12)
    remember_me:bool=False

    @field_validator('username')
    @classmethod
    def validate_login_format(cls,value:str)->str:
        value=value.strip()
        if not re.fullmatch(r'[A-Za-z][A-Za-z0-9]*(?:\.[A-Za-z0-9]+)*',value):
            raise ValueError('O login deve começar com uma letra; use letras, números e pontos entre partes do nome.')
        return value

class PasswordRecoveryIn(BaseModel):
    username:str=Field(min_length=1,max_length=60)
    recovery_code:str=Field(min_length=16,max_length=128)
    new_password:str=Field(min_length=6,max_length=12)

    @field_validator('username')
    @classmethod
    def validate_login_format(cls,value:str)->str:
        value=value.strip()
        if not re.fullmatch(r'[A-Za-z][A-Za-z0-9]*(?:\.[A-Za-z0-9]+)*',value):
            raise ValueError('O login deve começar com uma letra; use letras, números e pontos entre partes do nome.')
        return value

def _set_login_cookie(response:Response,user:AppUser,remember_me:bool=False):
    duration=security.REMEMBER_SESSION_SECONDS if remember_me else security.SESSION_SECONDS
    token,csrf=security.create_session(user.id,duration)
    response.set_cookie('lmf_session',token,max_age=duration,httponly=True,samesite='strict',secure=False,path='/')
    return {'authenticated':True,'username':user.username,'csrf_token':csrf,'remember_me':remember_me}

@app.get('/api/auth/status')
def auth_status(request:Request,db:Session=Depends(get_db)):
    configured=db.query(AppUser).filter(AppUser.active==True).count()>0
    session=security.read_session(request.cookies.get('lmf_session'))
    user=db.get(AppUser,session['uid']) if session else None
    return {'setup_required':not configured,'authenticated':bool(user and user.active),'username':user.username if user and user.active else None,'csrf_token':session['csrf'] if user and user.active else None,'recovery_file':str(security.RECOVERY_CODE_PATH)}

@app.post('/api/auth/setup')
def auth_setup(payload:AuthIn,response:Response,db:Session=Depends(get_db)):
    if db.query(AppUser).count(): raise HTTPException(409,'A configuração inicial já foi concluída.')
    try: recovery_code=security.RECOVERY_CODE_PATH.read_text(encoding='utf-8').strip()
    except OSError: recovery_code=''
    if len(recovery_code)<16: recovery_code=security.new_recovery_code()
    user=AppUser(username=payload.username,password_hash=security.password_hash(payload.password),recovery_code_hash=security.password_hash(recovery_code),is_admin=True,active=True)
    db.add(user); db.commit(); db.refresh(user)
    security.save_recovery_code(recovery_code)
    result=_set_login_cookie(response,user,payload.remember_me)
    result['recovery_file']=str(security.RECOVERY_CODE_PATH)
    result['recovery_code']=recovery_code
    return result

@app.post('/api/auth/login')
def auth_login(payload:AuthIn,request:Request,response:Response,db:Session=Depends(get_db)):
    address=request.client.host if request.client else 'local'
    now=time_module.time(); attempts=[t for t in _login_attempts.get(address,[]) if now-t<300]
    if len(attempts)>=8: raise HTTPException(429,'Muitas tentativas. Aguarde cinco minutos e tente novamente.')
    user=db.query(AppUser).filter(AppUser.username==payload.username,AppUser.active==True).first()
    if not user or not security.verify_password(payload.password,user.password_hash):
        attempts.append(now); _login_attempts[address]=attempts
        raise HTTPException(401,'Usuário ou senha inválidos.')
    _login_attempts.pop(address,None)
    return _set_login_cookie(response,user,payload.remember_me)

@app.post('/api/auth/forgot-password')
def auth_forgot_password(payload:PasswordRecoveryIn,request:Request,db:Session=Depends(get_db)):
    address=request.client.host if request.client else 'local'
    now=time_module.time()
    attempts=[stamp for stamp in _recovery_attempts.get(address,[]) if now-stamp<900]
    if len(attempts)>=5:
        raise HTTPException(429,'Muitas tentativas de recuperação. Aguarde 15 minutos.')
    attempts.append(now); _recovery_attempts[address]=attempts
    user=db.query(AppUser).filter(AppUser.username==payload.username,AppUser.active==True,AppUser.is_admin==True).first()
    if not user or not user.recovery_code_hash or not security.verify_password(payload.recovery_code,user.recovery_code_hash):
        raise HTTPException(400,'Login ou código de recuperação inválido.')
    new_code=security.new_recovery_code()
    old_password_hash=user.password_hash
    old_recovery_hash=user.recovery_code_hash
    user.password_hash=security.password_hash(payload.new_password)
    user.recovery_code_hash=security.password_hash(new_code)
    db.commit()
    try:
        security.save_recovery_code(new_code)
        security.rotate_session_secret()
    except OSError:
        user.password_hash=old_password_hash
        user.recovery_code_hash=old_recovery_hash
        db.commit()
        raise HTTPException(500,'Não foi possível salvar a recuperação. O arquivo de código não foi atualizado.')
    _recovery_attempts.pop(address,None)
    return {'status':'senha redefinida','message':'Senha redefinida. O novo código foi salvo no arquivo local de recuperação.','recovery_file':str(security.RECOVERY_CODE_PATH)}

class OperationAuthorizationIn(BaseModel):
    operation:str=Field(min_length=1,max_length=80)
    target:str=Field(min_length=1,max_length=120)
    password:SecretStr=Field(min_length=1,max_length=200)

@app.post('/api/auth/authorize-operation')
def auth_authorize_operation(payload:OperationAuthorizationIn,request:Request,response:Response,db:Session=Depends(get_db)):
    """Valida a senha do administrador (no backend) e emite um token de uso único para UMA operação sobre UM alvo."""
    response.headers['Cache-Control']='no-store'
    return authorization.issue_authorization(db,request,payload.operation,payload.target,payload.password.get_secret_value())

@app.get('/api/audit')
def audit_history(request:Request,limit:int=Query(100,ge=1,le=500),outcome:str|None=None,db:Session=Depends(get_db)):
    if not getattr(request.state,'is_admin',False): raise HTTPException(403,'Apenas o administrador pode consultar a auditoria.')
    return authorization.list_audit(db,limit,outcome)

@app.post('/api/auth/logout')
def auth_logout(response:Response):
    response.delete_cookie('lmf_session',path='/')
    return {'authenticated':False}

class ProductIn(BaseModel):
    model_config=ConfigDict(from_attributes=True)
    id:int|None=None
    sku:str|None=None; internal_code:str|None=None; name:str=Field(min_length=1); category:str|None=None; supplier:str|None=None; supplier_id:int|None=None; brand:str|None=None; model:str|None=None
    quantity:int=Field(default=0,ge=0); min_quantity:int=Field(default=0,ge=0); cost:float=Field(default=0,ge=0); price:float=Field(default=0,ge=0); active:bool=True; notes:str|None=None
class CustomerIn(BaseModel):
    model_config=ConfigDict(from_attributes=True)
    id:int|None=None
    name:str=Field(min_length=1); phone:str|None=None; whatsapp:str|None=None; email:str|None=None; document:str|None=None; address:str|None=None; notes:str|None=None; active:bool=True; public_seq:int|None=None
class SupplierIn(BaseModel):
    model_config=ConfigDict(from_attributes=True)
    id:int|None=None
    name:str=Field(min_length=1); contact:str|None=None; phone:str|None=None; whatsapp:str|None=None; email:str|None=None; website:str|None=None; categories:str|None=None; lead_time_days:int|None=None; active:bool=True; notes:str|None=None
class CustomerInput(CustomerIn):
    """Somente entrada: valida telefone/WhatsApp (DDD + até 9 dígitos) e CPF (11 dígitos)."""
    @field_validator('phone')
    @classmethod
    def _check_phone(cls,v): return normalize_phone(v,'Telefone')
    @field_validator('whatsapp')
    @classmethod
    def _check_whatsapp(cls,v): return normalize_phone(v,'WhatsApp')
    @field_validator('document')
    @classmethod
    def _check_document(cls,v): return normalize_cpf(v)
class SupplierInput(SupplierIn):
    @field_validator('phone')
    @classmethod
    def _check_phone(cls,v): return normalize_phone(v,'Telefone')
    @field_validator('whatsapp')
    @classmethod
    def _check_whatsapp(cls,v): return normalize_phone(v,'WhatsApp')
class ServiceIn(BaseModel):
    model_config=ConfigDict(from_attributes=True)
    id:int|None=None; number:str|None=None
    customer_id:int; equipment:str=Field(min_length=1); serial:str|None=None; equipment_type:str|None=None; reported_fault:str=Field(min_length=1); diagnosis:str|None=None
    status:str='Recebido'; estimate:float=Field(default=0,ge=0); cost:float=Field(default=0,ge=0); charged:float=Field(default=0,ge=0); approval:str='Pendente'; entered_on:date=Field(default_factory=date.today); due_on:date|None=None; delivered_on:date|None=None; completed_on:date|None=None; accessories:str|None=None; physical_condition:str|None=None; service_executed:str|None=None; notes:str|None=None
class NewCustomerData(BaseModel):
    name:str|None=None; document:str|None=None; phone:str|None=None; whatsapp:str|None=None; email:str|None=None
    @field_validator('phone')
    @classmethod
    def _p(cls,v): return normalize_phone(v,'Telefone')
    @field_validator('whatsapp')
    @classmethod
    def _w(cls,v): return normalize_phone(v,'WhatsApp')
    @field_validator('document')
    @classmethod
    def _d(cls,v): return normalize_cpf(v)
class ServiceInput(ServiceIn):
    """Entrada da OS: cliente existente (customer_id) OU dados do cliente (CPF) para localizar/cadastrar."""
    customer_id:int|None=None
    customer:NewCustomerData|None=None
    equipment_type:Literal['Celular','Computador','Notebook','Videogame']
    serial:str
    @field_validator('serial',mode='before')
    @classmethod
    def _serial(cls,v): return normalize_serial(v)
class OrderIn(BaseModel):
    model_config=ConfigDict(from_attributes=True)
    id:int|None=None; number:str|None=None
    kind:str='cliente'; customer_id:int|None=None; supplier_id:int|None=None; description:str=Field(min_length=1); quantity:int=Field(default=1,gt=0)
    total:float=Field(default=0,ge=0); paid:float=Field(default=0,ge=0); status:str='Solicitação recebida'; due_on:date|None=None; received:bool=False; notes:str|None=None
class GuaranteeIn(BaseModel):
    model_config=ConfigDict(from_attributes=True)
    id:int|None=None
    customer_id:int; service_order_id:int|None=None; item:str=Field(min_length=1); guarantee_type:str='serviço'; starts_on:date; ends_on:date; duration_days:int|None=None; status:str='ativa'; conditions:str|None=None; notes:str|None=None
class GuaranteeInput(GuaranteeIn):
    ends_on:date|None=None  # o vencimento é calculado pelo servidor (início + 90 dias)
class GuaranteeOccurrenceIn(BaseModel):
    issue:str=Field(min_length=2); resolution:str|None=None; status:Literal['aberta','em análise','resolvida']='aberta'
class FinanceIn(BaseModel):
    model_config=ConfigDict(from_attributes=True)
    id:int|None=None
    kind:Literal['receita','despesa']; category:str; description:str; amount:float=Field(gt=0); status:str='pendente'; competence_on:date|None=None; due_on:date|None=None; paid_on:date|None=None; payment_method:str|None=None; bank_account_id:int|None=None; customer_id:int|None=None; service_order_id:int|None=None; order_id:int|None=None; sale_id:int|None=None; external_id:str|None=None; reference:str|None=None
class BankAccountIn(BaseModel):
    model_config=ConfigDict(from_attributes=True)
    id:int|None=None
    name:str=Field(min_length=1); business_unit:str='Assistência'; provider:str='InfinitePay'; opening_balance:float=0; opening_balance_date:date|None=None; active:bool=True; notes:str|None=None
class RemoveOpeningBalanceIn(BaseModel):
    account_name_confirmation:str=Field(min_length=1,max_length=120)
    confirm_linked_transactions:bool=False
class SaleLineIn(BaseModel):
    product_id:int; quantity:int=Field(gt=0); unit_price:float=Field(ge=0); discount:float=Field(default=0,ge=0)
class SaleIn(BaseModel):
    customer_id:int|None=None; items:list[SaleLineIn]=Field(min_length=1); discount:float=Field(default=0,ge=0)
    received_amount:float=Field(default=0,ge=0); payment_method:str|None=None; bank_account_id:int|None=None; status:Literal['confirmada','pendente']='confirmada'; notes:str|None=None; request_key:str|None=None
class SaleEditIn(BaseModel):
    customer_id:int|None=None
    payment_method:str|None=None
    bank_account_id:int|None=None
    notes:str|None=None
class PurchaseReceiptIn(BaseModel):
    product_id:int; quantity:int=Field(gt=0); unit_cost:float=Field(default=0,ge=0)

def crud_routes(path, model, schema, searchable=(), input_schema=None, delete_guard=None, prepare=None, before_list=None):
    in_schema=input_schema or schema  # entrada pode ser mais rígida que a resposta
    def listing(q:str|None=None, db:Session=Depends(get_db)):
        if before_list: before_list(db)
        query=db.query(model)
        if q and searchable: query=query.filter(or_(*[getattr(model,c).ilike(f'%{q}%') for c in searchable]))
        return query.order_by(model.id.desc()).limit(500).all()
    listing.__name__=f'list_{path.replace("-","_")}'
    app.get(f'/api/{path}',response_model=list[schema])(listing)
    def create(payload:in_schema,db:Session=Depends(get_db)):
        values=payload.model_dump(exclude={'id','number','public_seq'})
        if prepare: prepare(values,None,db)
        if model is Order:
            year=date.today().year; prefix=f'LMF-PED-{year}-'
            last=db.query(Order).filter(Order.number.like(prefix+'%')).order_by(Order.id.desc()).first()
            seq=(int(last.number.rsplit('-',1)[-1])+1) if last else 1
            values['number']=f'{prefix}{seq:05d}'
        obj=model(**values); db.add(obj)
        try: db.commit(); db.refresh(obj)
        except Exception as e: db.rollback(); raise HTTPException(409,f'Não foi possível salvar: {e}')
        return obj
    create.__name__=f'create_{path.replace("-","_")}'
    app.post(f'/api/{path}',response_model=schema)(create)
    def update(item_id:int,payload:in_schema,db:Session=Depends(get_db)):
        obj=db.get(model,item_id)
        if not obj: raise HTTPException(404,'Registro não encontrado')
        values=payload.model_dump(exclude={'id','number','public_seq'})
        if prepare: prepare(values,obj,db)
        for k,v in values.items(): setattr(obj,k,v)
        db.commit(); db.refresh(obj); return obj
    update.__name__=f'update_{path.replace("-","_")}'
    app.put(f'/api/{path}/{{item_id}}',response_model=schema)(update)
    def delete(item_id:int,request:Request,confirm:bool=Query(False),db:Session=Depends(get_db)):
        if not confirm: raise HTTPException(400,'Confirme a exclusão do registro.')
        if delete_guard: delete_guard(db,request,item_id)  # autorização administrativa, validada no backend
        obj=db.get(model,item_id)
        if not obj: raise HTTPException(404,'Registro não encontrado')
        db.delete(obj)
        try: db.commit()
        except Exception:
            db.rollback()
            raise HTTPException(409,'Este registro possui movimentações ou vínculos. Remova ou ajuste os registros dependentes antes de excluí-lo.')
        return {'deleted':True,'id':item_id}
    delete.__name__=f'delete_{path.replace("-","_")}'
    app.delete(f'/api/{path}/{{item_id}}')(delete)

crud_routes('products',Product,ProductIn,('sku','internal_code','name','category','supplier','brand','model','notes'))
@app.get('/api/customers/lookup')
def lookup_customer(document:str,db:Session=Depends(get_db)):
    try: cpf=normalize_cpf(document)
    except ValueError as e: raise HTTPException(400,str(e))
    found=customer_rules.find_by_document(db,cpf)
    if not found: raise HTTPException(404,'Cliente não encontrado')
    return {'id':found.id,'public_seq':found.public_seq,'public_code':customer_rules.public_code(found.public_seq),'name':found.name,'document':found.document,'phone':found.phone,'whatsapp':found.whatsapp,'email':found.email}
crud_routes('customers',Customer,CustomerIn,('name','phone','document','email'),input_schema=CustomerInput,prepare=customer_rules.prepare_customer)
crud_routes('suppliers',Supplier,SupplierIn,('name','categories','contact','phone','whatsapp','email','website','notes'),input_schema=SupplierInput)
crud_routes('orders',Order,OrderIn,('number','description','status'))
crud_routes('guarantees',Guarantee,GuaranteeIn,('item','conditions'),input_schema=GuaranteeInput,prepare=guarantee_rules.prepare_guarantee,before_list=guarantee_rules.expire_due_guarantees)
crud_routes('finance',FinancialEntry,FinanceIn,('description','category','kind','status','reference','payment_method','external_id'),delete_guard=lambda db,request,item_id:authorization.require_authorization(db,request,'financeiro.lancamento.excluir',item_id))

@app.get('/api/accounts')
def list_accounts(db:Session=Depends(get_db)):
    results=[]
    for account in db.query(BankAccount).order_by(BankAccount.id).all():
        income=db.query(func.sum(FinancialEntry.amount)).filter(FinancialEntry.bank_account_id==account.id,FinancialEntry.kind=='receita',FinancialEntry.paid_on.isnot(None)).scalar() or 0
        expenses=db.query(func.sum(FinancialEntry.amount)).filter(FinancialEntry.bank_account_id==account.id,FinancialEntry.kind=='despesa',FinancialEntry.paid_on.isnot(None)).scalar() or 0
        linked=_account_linked_transaction_count(db,account.id)
        results.append({'id':account.id,'name':account.name,'business_unit':account.business_unit,'provider':account.provider,'opening_balance':account.opening_balance,'opening_balance_date':account.opening_balance_date,'balance':account.opening_balance+income-expenses,'active':account.active,'notes':account.notes,'linked_transaction_count':linked})
    return results

def _account_linked_transaction_count(db:Session,account_id:int)->int:
    return db.query(FinancialEntry).filter_by(bank_account_id=account_id).count()+db.query(BankTransaction).filter_by(bank_account_id=account_id).count()

@app.post('/api/accounts')
def create_account(payload:BankAccountIn,db:Session=Depends(get_db)):
    values=payload.model_dump(exclude={'id'})
    if values['opening_balance'] and values.get('opening_balance_date') is None: values['opening_balance_date']=date.today()
    account=BankAccount(**values); db.add(account); db.commit(); db.refresh(account); return account
@app.put('/api/accounts/{item_id}')
def update_account(item_id:int,payload:BankAccountIn,request:Request,db:Session=Depends(get_db)):
    account=db.get(BankAccount,item_id)
    if not account: raise HTTPException(404,'Conta não encontrada')
    values=payload.model_dump(exclude={'id'})
    old_balance=account.opening_balance; old_date=account.opening_balance_date
    new_date=values['opening_balance_date'] if 'opening_balance_date' in payload.model_fields_set else old_date
    if values['opening_balance'] and new_date is None and values['opening_balance']!=old_balance: new_date=date.today()
    values['opening_balance_date']=new_date
    if old_balance!=values['opening_balance'] or old_date!=new_date: authorization.require_authorization(db,request,'financeiro.conta.saldo_inicial.alterar',item_id)
    for key,value in values.items(): setattr(account,key,value)
    if old_balance!=account.opening_balance or old_date!=account.opening_balance_date:
        db.add(AccountBalanceAdjustment(bank_account_id=account.id,user_id=request.state.user_id,old_balance=old_balance,new_balance=account.opening_balance,old_reference_date=old_date,new_reference_date=account.opening_balance_date,action='alterado',reason='Alteração administrativa do saldo inicial'))
    db.commit(); db.refresh(account); return account

@app.delete('/api/accounts/{item_id}')
def delete_account(item_id:int,request:Request,confirm:bool=Query(False),db:Session=Depends(get_db)):
    if not confirm: raise HTTPException(400,'Confirme a exclusão da conta.')
    authorization.require_authorization(db,request,'financeiro.conta.excluir',item_id)
    account=db.get(BankAccount,item_id)
    if not account: raise HTTPException(404,'Conta não encontrada')
    if _account_linked_transaction_count(db,item_id) or db.query(AccountBalanceAdjustment).filter_by(bank_account_id=item_id).count():
        raise HTTPException(409,'A conta possui movimentações ou ajustes vinculados. Edite ou desative a conta para preservar o histórico financeiro.')
    db.delete(account); db.commit(); return {'deleted':True,'id':item_id}

@app.post('/api/accounts/{item_id}/remove-opening-balance')
def remove_opening_balance(item_id:int,payload:RemoveOpeningBalanceIn,request:Request,db:Session=Depends(get_db)):
    if not getattr(request.state,'is_admin',False): raise HTTPException(403,'Apenas o administrador pode remover saldo inicial.')
    account=db.query(BankAccount).filter_by(id=item_id).with_for_update().first()
    if not account: raise HTTPException(404,'Conta não encontrada')
    if account.opening_balance==0:
        return {'status':'sem saldo inicial','balance':_computed_account_balance(db,account)}
    if payload.account_name_confirmation.strip()!=account.name: raise HTTPException(400,'Confirmação do nome da conta não corresponde.')
    linked=_account_linked_transaction_count(db,account.id)
    if linked and not payload.confirm_linked_transactions: raise HTTPException(409,f'Esta conta tem {linked} movimentação(ões) vinculada(s). Confirme essa etapa adicional para continuar.')
    authorization.require_authorization(db,request,'financeiro.conta.saldo_inicial.remover',item_id)
    old_balance=account.opening_balance; old_date=account.opening_balance_date
    account.opening_balance=0; account.opening_balance_date=None
    db.add(AccountBalanceAdjustment(bank_account_id=account.id,user_id=request.state.user_id,old_balance=old_balance,new_balance=0,old_reference_date=old_date,new_reference_date=None,action='saldo inicial removido',reason='Remoção administrativa; movimentações históricas preservadas'))
    db.commit(); db.refresh(account)
    return {'status':'saldo inicial removido','opening_balance':account.opening_balance,'balance':_computed_account_balance(db,account),'linked_transaction_count':linked}

def _computed_account_balance(db:Session,account:BankAccount)->float:
    income=db.query(func.sum(FinancialEntry.amount)).filter(FinancialEntry.bank_account_id==account.id,FinancialEntry.kind=='receita',FinancialEntry.paid_on.isnot(None)).scalar() or 0
    expenses=db.query(func.sum(FinancialEntry.amount)).filter(FinancialEntry.bank_account_id==account.id,FinancialEntry.kind=='despesa',FinancialEntry.paid_on.isnot(None)).scalar() or 0
    return account.opening_balance+income-expenses

@app.get('/api/sales')
def list_sales(db:Session=Depends(get_db)):
    result=[]
    for sale in db.query(Sale).order_by(Sale.id.desc()).limit(1000):
        items=db.query(SaleItem).filter_by(sale_id=sale.id).all()
        result.append({'id':sale.id,'number':sale.number,'customer_id':sale.customer_id,'sold_at':sale.sold_at,'status':sale.status,'total':sale.total,'discount':sale.discount,'estimated_cost':sale.estimated_cost,'received_amount':sale.received_amount,'payment_method':sale.payment_method,'items':[{'product_id':i.product_id,'product':(db.get(Product,i.product_id).name if db.get(Product,i.product_id) else 'Produto arquivado'),'quantity':i.quantity,'unit_price':i.unit_price,'discount':i.discount} for i in items]})
    return result

@app.post('/api/sales')
def create_sale(payload:SaleIn,idempotency_key:str|None=Header(default=None,alias='Idempotency-Key'),db:Session=Depends(get_db)):
    request_key=idempotency_key or payload.request_key
    if request_key:
        existing=db.query(Sale).filter_by(request_key=request_key).first()
        if existing:
            old_lines=db.query(SaleItem).filter_by(sale_id=existing.id).order_by(SaleItem.id).all()
            same=(existing.customer_id==payload.customer_id and existing.discount==payload.discount and existing.received_amount==payload.received_amount and existing.payment_method==payload.payment_method and existing.bank_account_id==payload.bank_account_id and len(old_lines)==len(payload.items) and all((a.product_id,a.quantity,a.unit_price,a.discount)==(b.product_id,b.quantity,b.unit_price,b.discount) for a,b in zip(old_lines,payload.items)))
            if not same: raise HTTPException(409,'Idempotency-Key já foi utilizado com dados diferentes.')
            return {'id':existing.id,'number':existing.number,'status':existing.status,'total':existing.total,'estimated_cost':existing.estimated_cost,'received_amount':existing.received_amount,'duplicate_request':True}
    if payload.customer_id and not db.get(Customer,payload.customer_id): raise HTTPException(400,'Cliente não encontrado')
    if payload.bank_account_id and not db.get(BankAccount,payload.bank_account_id): raise HTTPException(400,'Conta bancária não encontrada')
    try:
        line_data=[]; total=0.0; estimated_cost=0.0
        for line in payload.items:
            product=db.query(Product).filter_by(id=line.product_id,active=True).with_for_update().first()
            if not product: raise HTTPException(400,f'Produto {line.product_id} não encontrado ou inativo')
            line_total=line.quantity*line.unit_price-line.discount
            if line.discount>line.quantity*line.unit_price: raise HTTPException(400,'Desconto do item maior que seu valor')
            if payload.status=='confirmada' and product.quantity<line.quantity: raise HTTPException(409,f'Estoque insuficiente para {product.name}')
            line_data.append((line,product,line_total)); total+=line_total; estimated_cost+=line.quantity*product.cost
        if payload.discount>total: raise HTTPException(400,'Desconto total maior que o valor da venda')
        total-=payload.discount
        if payload.received_amount>total: raise HTTPException(400,'Valor recebido não pode superar o total da venda')
        year=date.today().year; prefix=f'LMF-VEN-{year}-'
        last=db.query(Sale).filter(Sale.number.like(prefix+'%')).order_by(Sale.id.desc()).first()
        seq=(int(last.number.rsplit('-',1)[-1])+1) if last else 1
        sale=Sale(number=f'{prefix}{seq:05d}',request_key=request_key,customer_id=payload.customer_id,status=payload.status,total=total,discount=payload.discount,estimated_cost=estimated_cost,received_amount=payload.received_amount,payment_method=payload.payment_method,bank_account_id=payload.bank_account_id,notes=payload.notes)
        db.add(sale); db.flush()
        for index,(line,product,line_total) in enumerate(line_data,1):
            db.add(SaleItem(sale_id=sale.id,product_id=product.id,quantity=line.quantity,unit_price=line.unit_price,unit_cost=product.cost,discount=line.discount))
            if payload.status=='confirmada':
                product.quantity-=line.quantity
                db.add(Movement(product_id=product.id,kind='saida',quantity=line.quantity,reason='Venda confirmada',reference=sale.number,source_key=f'sale:{sale.id}:line:{index}'))
        if payload.received_amount>0:
            db.add(FinancialEntry(kind='receita',category='Venda de produto',description=f'Recebimento da venda {sale.number}',amount=payload.received_amount,status='recebido',competence_on=date.today(),due_on=date.today(),paid_on=date.today(),payment_method=payload.payment_method,bank_account_id=payload.bank_account_id,customer_id=payload.customer_id,sale_id=sale.id,external_id=f'sale-receipt:{sale.id}',reference=f'SALE:{sale.number}:received'))
        outstanding=round(total-payload.received_amount,2)
        if outstanding>0:
            db.add(FinancialEntry(kind='receita',category='Venda a receber',description=f'Saldo pendente da venda {sale.number}',amount=outstanding,status='pendente',competence_on=date.today(),due_on=date.today(),payment_method=payload.payment_method,bank_account_id=payload.bank_account_id,customer_id=payload.customer_id,sale_id=sale.id,external_id=f'sale-receivable:{sale.id}',reference=f'SALE:{sale.number}:receivable'))
        db.commit(); db.refresh(sale)
        return {'id':sale.id,'number':sale.number,'status':sale.status,'total':sale.total,'estimated_cost':sale.estimated_cost,'received_amount':sale.received_amount}
    except HTTPException: db.rollback(); raise
    except Exception as e: db.rollback(); raise HTTPException(409,f'Venda não registrada: {e}')

@app.post('/api/sales/{sale_id}/cancel')
def cancel_sale(sale_id:int,request:Request,db:Session=Depends(get_db)):
    sale=db.query(Sale).filter_by(id=sale_id).with_for_update().first()
    if not sale: raise HTTPException(404,'Venda não encontrada')
    authorization.require_authorization(db,request,'financeiro.venda.cancelar',sale_id)
    if sale.status=='cancelada': return {'status':'cancelada','duplicate_request':True}
    try:
        if sale.status=='confirmada':
            items=db.query(SaleItem).filter_by(sale_id=sale.id).order_by(SaleItem.id).all()
            for index,item in enumerate(items,1):
                product=db.query(Product).filter_by(id=item.product_id).with_for_update().first()
                if product:
                    product.quantity+=item.quantity
                    db.add(Movement(product_id=product.id,kind='entrada',quantity=item.quantity,reason=f'Estorno da venda {sale.number}',reference=sale.number,source_key=f'sale-cancel:{sale.id}:line:{index}'))
        if sale.received_amount>0 and not db.query(FinancialEntry).filter_by(external_id=f'sale-refund:{sale.id}').first():
            db.add(FinancialEntry(kind='despesa',category='Estorno de venda',description=f'Estorno da venda {sale.number}',amount=sale.received_amount,status='pago',competence_on=date.today(),due_on=date.today(),paid_on=date.today(),payment_method=sale.payment_method,bank_account_id=sale.bank_account_id,customer_id=sale.customer_id,sale_id=sale.id,external_id=f'sale-refund:{sale.id}',reference=f'SALE:{sale.number}:refund'))
        sale.status='cancelada'; db.commit(); return {'status':'cancelada','refunded':sale.received_amount}
    except Exception as e: db.rollback(); raise HTTPException(409,f'Falha ao cancelar venda: {e}')

@app.put('/api/sales/{sale_id}')
def update_sale(sale_id:int,payload:SaleEditIn,db:Session=Depends(get_db)):
    sale=db.get(Sale,sale_id)
    if not sale: raise HTTPException(404,'Venda não encontrada')
    if sale.status=='cancelada': raise HTTPException(409,'Uma venda cancelada não pode ser editada.')
    for key,value in payload.model_dump().items(): setattr(sale,key,value)
    for entry in db.query(FinancialEntry).filter_by(sale_id=sale_id).all():
        entry.customer_id=sale.customer_id; entry.bank_account_id=sale.bank_account_id; entry.payment_method=sale.payment_method
    db.commit(); return {'updated':True,'id':sale_id}

@app.delete('/api/sales/{sale_id}')
def delete_sale(sale_id:int,request:Request,confirm:bool=Query(False),db:Session=Depends(get_db)):
    if not confirm: raise HTTPException(400,'Confirme a exclusão da venda.')
    authorization.require_authorization(db,request,'financeiro.venda.excluir',sale_id)
    sale=db.get(Sale,sale_id)
    if not sale: raise HTTPException(404,'Venda não encontrada')
    try:
        items=db.query(SaleItem).filter_by(sale_id=sale_id).order_by(SaleItem.id).all()
        if sale.status=='confirmada':
            for index,item in enumerate(items,1):
                product=db.get(Product,item.product_id)
                if product:
                    product.quantity+=item.quantity
                    db.add(Movement(product_id=product.id,kind='entrada',quantity=item.quantity,reason=f'Ajuste administrativo pela exclusão da venda {sale.number}',reference=sale.number,source_key=f'sale-delete:{sale.id}:line:{index}'))
        linked_entries=db.query(FinancialEntry).filter_by(sale_id=sale_id).all()
        linked_ids=[entry.id for entry in linked_entries]
        if linked_ids:
            db.query(BankTransaction).filter(BankTransaction.financial_entry_id.in_(linked_ids)).update({BankTransaction.financial_entry_id:None},synchronize_session=False)
            for entry in linked_entries: db.delete(entry)
        db.query(Movement).filter(Movement.reference==sale.number).delete(synchronize_session=False)
        for item in items: db.delete(item)
        db.delete(sale); db.commit(); return {'deleted':True,'id':sale_id,'stock_restored':sale.status=='confirmada'}
    except Exception as error:
        db.rollback(); raise HTTPException(409,f'Não foi possível excluir a venda: {error}')

@app.post('/api/service-orders/{service_id}/parts')
def use_service_part(service_id:int,product_id:int,quantity:int=Query(gt=0),db:Session=Depends(get_db)):
    service=db.get(ServiceOrder,service_id)
    if not service: raise HTTPException(404,'Ordem de serviço não encontrada')
    product=db.query(Product).filter_by(id=product_id).with_for_update().first()
    if not product: raise HTTPException(404,'Produto não encontrado')
    if product.quantity<quantity: raise HTTPException(409,'Estoque insuficiente')
    key=f'service:{service_id}:product:{product_id}:{datetime.now().strftime("%Y%m%d%H%M%S%f")}'
    product.quantity-=quantity
    db.add(ServicePart(service_order_id=service_id,product_id=product_id,quantity=quantity,unit_cost=product.cost,unit_price=product.price,source_key=key))
    db.add(Movement(product_id=product_id,kind='saida',quantity=quantity,reason=f'Peça utilizada na OS {service.number}',reference=service.number,source_key=key))
    db.commit(); return {'quantity':product.quantity,'source_key':key}

@app.post('/api/orders/{order_id}/receive')
def receive_purchase(order_id:int,payload:PurchaseReceiptIn,idempotency_key:str|None=Header(default=None,alias='Idempotency-Key'),db:Session=Depends(get_db)):
    if not idempotency_key: raise HTTPException(400,'Envie o cabeçalho Idempotency-Key para evitar entrada duplicada.')
    existing=db.query(PurchaseReceipt).filter_by(request_key=idempotency_key).first()
    if existing:
        if (existing.order_id,existing.product_id,existing.quantity)!=(order_id,payload.product_id,payload.quantity): raise HTTPException(409,'Idempotency-Key já foi usado com dados diferentes.')
        return {'status':'recebido','duplicate_request':True,'quantity':existing.quantity}
    order=db.query(Order).filter_by(id=order_id).with_for_update().first()
    if not order: raise HTTPException(404,'Pedido não encontrado')
    if order.kind not in ('compra','fornecedor'): raise HTTPException(400,'Somente pedidos de compra a fornecedor podem receber estoque.')
    if order.status in ('Cancelado','Concluído') or order.received: raise HTTPException(409,'Este pedido está encerrado.')
    item_rows=db.query(OrderItem).filter_by(order_id=order_id).all()
    if item_rows and any(item.product_id!=payload.product_id for item in item_rows): raise HTTPException(400,'Este fluxo recebe um produto por pedido; separe produtos em pedidos distintos.')
    item=next((x for x in item_rows if x.product_id==payload.product_id),None)
    if item is None:
        item=OrderItem(order_id=order_id,product_id=payload.product_id,quantity=order.quantity,unit_price=payload.unit_cost,received_quantity=0)
        db.add(item); db.flush()
    if item.received_quantity+payload.quantity>item.quantity: raise HTTPException(409,'Quantidade recebida excede a quantidade pedida.')
    product=db.query(Product).filter_by(id=payload.product_id).with_for_update().first()
    if not product: raise HTTPException(404,'Produto não encontrado')
    try:
        receipt=PurchaseReceipt(request_key=idempotency_key,order_id=order_id,order_item_id=item.id,product_id=product.id,quantity=payload.quantity,unit_cost=payload.unit_cost)
        db.add(receipt); db.flush()
        item.received_quantity+=payload.quantity; product.quantity+=payload.quantity
        db.add(Movement(product_id=product.id,kind='entrada',quantity=payload.quantity,reason=f'Recebimento do pedido {order.number}',reference=order.number,source_key=f'purchase-receipt:{receipt.id}'))
        if item.received_quantity==item.quantity: order.received=True; order.status='Concluído'
        db.commit(); return {'status':'recebido','quantity':payload.quantity,'received_total':item.received_quantity,'product_quantity':product.quantity,'order_complete':order.received}
    except Exception as e: db.rollback(); raise HTTPException(409,f'Falha ao receber o pedido: {e}')

@app.get('/',response_class=HTMLResponse)
def index(): return FileResponse(Path(__file__).parent/'templates'/'index.html')

@app.get('/api/dashboard')
def dashboard(db:Session=Depends(get_db)):
    today=date.today()
    recent_os=db.query(ServiceOrder).order_by(ServiceOrder.id.desc()).limit(5).all()
    os_rows=[{'number':o.number,'customer':(db.get(Customer,o.customer_id).name if db.get(Customer,o.customer_id) else '—'),'equipment':o.equipment,'status':o.status,'date':o.entered_on.isoformat() if o.entered_on else ''} for o in recent_os]
    low_products=db.query(Product).filter(Product.quantity<=Product.min_quantity,Product.active==True).order_by(Product.quantity.asc(),Product.name.asc()).limit(5).all()
    low_rows=[{'name':p.name,'category':p.category or '—','quantity':p.quantity} for p in low_products]
    pending=db.query(Order).filter(Order.status.notin_(['Concluído','Cancelado'])).order_by(Order.due_on.asc().nullslast(),Order.id.desc()).limit(5).all()
    pending_rows=[]
    for order in pending:
        relation=db.get(Customer,order.customer_id) if order.customer_id else (db.get(Supplier,order.supplier_id) if order.supplier_id else None)
        pending_rows.append({'number':order.number,'party':relation.name if relation else '—','description':order.description,'due_on':order.due_on.isoformat() if order.due_on else '—','status':order.status})
    return {'products':db.query(Product).filter(Product.active==True).count(),'low_stock':db.query(Product).filter(Product.quantity<=Product.min_quantity,Product.active==True).count(),
      'stock_cost':db.query(func.sum(Product.quantity*Product.cost)).scalar() or 0,'service_in_progress':db.query(ServiceOrder).filter(ServiceOrder.status.notin_(['Entregue','Cancelado'])).count(),
      'awaiting_approval':db.query(ServiceOrder).filter(ServiceOrder.status=='Aguardando aprovação').count(),'orders_pending':db.query(Order).filter(Order.status.notin_(['Concluído','Cancelado'])).count(),
      'orders_late':db.query(Order).filter(Order.due_on<today,Order.status.notin_(['Concluído','Cancelado'])).count(),'guarantees_expiring':db.query(Guarantee).filter(Guarantee.ends_on>=today,Guarantee.ends_on<=date.fromordinal(today.toordinal()+30)).count(),
      'suppliers':db.query(Supplier).filter(Supplier.active==True).count(),'income':db.query(func.sum(FinancialEntry.amount)).filter(FinancialEntry.kind=='receita',FinancialEntry.paid_on.isnot(None),func.strftime('%Y-%m',FinancialEntry.paid_on)==today.strftime('%Y-%m')).scalar() or 0,
      'expenses':db.query(func.sum(FinancialEntry.amount)).filter(FinancialEntry.kind=='despesa',FinancialEntry.paid_on.isnot(None),func.strftime('%Y-%m',FinancialEntry.paid_on)==today.strftime('%Y-%m')).scalar() or 0,
      'recent_service_orders':os_rows,'low_stock_products':low_rows,'pending_order_rows':pending_rows,'as_of':today.isoformat()}

@app.get('/api/customers/{customer_id}/history')
def customer_history(customer_id:int,db:Session=Depends(get_db)):
    if not db.get(Customer,customer_id): raise HTTPException(404,'Cliente não encontrado')
    return {'service_orders':db.query(ServiceOrder).filter_by(customer_id=customer_id).all(),'orders':db.query(Order).filter_by(customer_id=customer_id).all(),'guarantees':db.query(Guarantee).filter_by(customer_id=customer_id).all(),'sales':db.query(Sale).filter_by(customer_id=customer_id).all(),'financial_entries':db.query(FinancialEntry).filter_by(customer_id=customer_id).all()}

@app.get('/api/guarantees/{guarantee_id}/occurrences')
def list_guarantee_occurrences(guarantee_id:int,db:Session=Depends(get_db)):
    if not db.get(Guarantee,guarantee_id): raise HTTPException(404,'Garantia não encontrada')
    return db.query(GuaranteeOccurrence).filter_by(guarantee_id=guarantee_id).order_by(GuaranteeOccurrence.id.desc()).all()

@app.post('/api/guarantees/{guarantee_id}/occurrences')
def create_guarantee_occurrence(guarantee_id:int,payload:GuaranteeOccurrenceIn,db:Session=Depends(get_db)):
    guarantee=db.get(Guarantee,guarantee_id)
    if not guarantee: raise HTTPException(404,'Garantia não encontrada')
    if guarantee.status!='ativa' or guarantee.ends_on<date.today(): raise HTTPException(409,'A garantia não está ativa')
    item=GuaranteeOccurrence(guarantee_id=guarantee_id,**payload.model_dump())
    if item.status=='resolvida': item.closed_on=date.today()
    db.add(item); db.commit(); db.refresh(item); return item

def _resolve_service_customer(db,payload):
    if payload.customer_id is not None:
        found=db.get(Customer,payload.customer_id)
        if not found: raise HTTPException(400,'Cliente não encontrado')
        return found
    if payload.customer is None: raise HTTPException(400,'Informe o cliente (CPF) da ordem de serviço.')
    return customer_rules.get_or_create_for_order(db,payload.customer.model_dump())

@app.post('/api/service-orders',response_model=ServiceIn)
def create_service(payload:ServiceInput,db:Session=Depends(get_db)):
    customer=_resolve_service_customer(db,payload)
    year=date.today().year
    last=db.query(ServiceOrder).filter(ServiceOrder.number.like(f'LMF-OS-{year}-%')).order_by(ServiceOrder.id.desc()).first()
    seq=int(last.number[-4:])+1 if last else 1
    values=payload.model_dump(exclude={'id','number','customer'}); values['customer_id']=customer.id
    obj=ServiceOrder(**values,number=f'LMF-OS-{year}-{seq:04d}'); db.add(obj); db.commit(); db.refresh(obj); return obj

@app.get('/api/service-orders',response_model=list[ServiceIn])
def list_services(q:str|None=None,equipment_type:str|None=None,db:Session=Depends(get_db)):
    query=db.query(ServiceOrder)
    if equipment_type: query=query.filter(ServiceOrder.equipment_type==equipment_type)
    if q:
        pattern=f'%{q}%'
        query=query.outerjoin(Customer,ServiceOrder.customer_id==Customer.id).filter(or_(ServiceOrder.number.ilike(pattern),ServiceOrder.equipment.ilike(pattern),ServiceOrder.serial.ilike(pattern),ServiceOrder.equipment_type.ilike(pattern),ServiceOrder.reported_fault.ilike(pattern),ServiceOrder.diagnosis.ilike(pattern),ServiceOrder.service_executed.ilike(pattern),ServiceOrder.status.ilike(pattern),ServiceOrder.approval.ilike(pattern),Customer.name.ilike(pattern)))
    return query.order_by(ServiceOrder.id.desc()).limit(500).all()

@app.put('/api/service-orders/{item_id}',response_model=ServiceIn)
def update_service(item_id:int,payload:ServiceInput,db:Session=Depends(get_db)):
    obj=db.get(ServiceOrder,item_id)
    if not obj: raise HTTPException(404,'Ordem de serviço não encontrada')
    if payload.status=='Entregue' and not payload.diagnosis: raise HTTPException(400,'Informe o diagnóstico antes de concluir a ordem')
    old_status=obj.status
    values=payload.model_dump(exclude={'id','number','customer'}); values['customer_id']=_resolve_service_customer(db,payload).id
    for key,value in values.items(): setattr(obj,key,value)
    if obj.status=='Entregue' and not obj.delivered_on: obj.delivered_on=date.today()
    if old_status!=obj.status: db.add(ServiceOrderEvent(service_order_id=obj.id,old_status=old_status,new_status=obj.status,details='Atualização operacional'))
    db.commit(); db.refresh(obj); return obj

@app.delete('/api/service-orders/{item_id}')
def delete_service(item_id:int,confirm:bool=Query(False),db:Session=Depends(get_db)):
    if not confirm: raise HTTPException(400,'Confirme a exclusão da ordem de serviço.')
    obj=db.get(ServiceOrder,item_id)
    if not obj: raise HTTPException(404,'Ordem de serviço não encontrada')
    if db.query(ServicePart).filter_by(service_order_id=item_id).count() or db.query(FinancialEntry).filter_by(service_order_id=item_id).count() or db.query(Guarantee).filter_by(service_order_id=item_id).count():
        raise HTTPException(409,'Esta OS tem peças, lançamentos financeiros ou garantias vinculados. Edite os vínculos ou mantenha a OS para preservar o histórico.')
    db.query(ServiceOrderEvent).filter_by(service_order_id=item_id).delete(synchronize_session=False)
    photo_files=[_photo_dir(item_id)/row.stored_name for row in db.query(ServiceOrderPhoto).filter_by(service_order_id=item_id).all()]
    db.query(ServiceOrderPhoto).filter_by(service_order_id=item_id).delete(synchronize_session=False)
    db.delete(obj); db.commit()
    _remove_photo_files(item_id,photo_files)  # só depois do commit: se falhar, os arquivos continuam
    return {'deleted':True,'id':item_id}

# --- Fotos das ordens de serviço (máx. 6 por OS) -------------------------------------------
MAX_PHOTOS_PER_ORDER=6
MAX_PHOTO_BYTES=8*1024*1024
PHOTO_EXTENSIONS={'image/jpeg':'.jpg','image/png':'.png','image/webp':'.webp'}
_STORED_NAME=re.compile(r'^[0-9a-f]{32}\.(jpg|png|webp)$')

def _photo_dir(order_id:int)->Path:
    return DATA_DIR/'documentos'/'ordens-servico'/str(int(order_id))

def _remove_photo_files(order_id:int,files)->None:
    for path in files:
        try: path.unlink(missing_ok=True)
        except OSError: _lmf_logger.warning('Não foi possível remover a foto %s',path)
    try: _photo_dir(order_id).rmdir()  # só remove se estiver vazia
    except OSError: pass

def _sniff_image(head:bytes)->str|None:
    """Tipo real pelo conteúdo (não confia na extensão nem no Content-Type enviados)."""
    if head.startswith(b'\xff\xd8\xff'): return 'image/jpeg'
    if head.startswith(b'\x89PNG\r\n\x1a\n'): return 'image/png'
    if head[:4]==b'RIFF' and head[8:12]==b'WEBP': return 'image/webp'
    return None

def _clean_filename(name:str|None)->str:
    base=Path((name or 'foto').replace('\\','/')).name
    base=''.join(ch for ch in base if ch.isprintable()).strip()
    return (base or 'foto')[:255]

def _photo_out(order_id:int,row:ServiceOrderPhoto)->dict:
    return {'id':row.id,'original_name':row.original_name,'content_type':row.content_type,'size':row.size,'url':f'/api/service-orders/{order_id}/photos/{row.id}/file'}

def _require_order(db:Session,order_id:int)->ServiceOrder:
    order=db.get(ServiceOrder,order_id)
    if not order: raise HTTPException(404,'Ordem de serviço não encontrada')
    return order

@app.get('/api/service-orders/{order_id}/photos')
def list_order_photos(order_id:int,db:Session=Depends(get_db)):
    _require_order(db,order_id)
    rows=db.query(ServiceOrderPhoto).filter_by(service_order_id=order_id).order_by(ServiceOrderPhoto.id).all()
    return {'max':MAX_PHOTOS_PER_ORDER,'items':[_photo_out(order_id,row) for row in rows]}

@app.post('/api/service-orders/{order_id}/photos')
def upload_order_photos(order_id:int,files:list[UploadFile]=File(...),db:Session=Depends(get_db)):
    _require_order(db,order_id)
    if not files: raise HTTPException(400,'Nenhuma foto enviada.')
    existing=db.query(func.count(ServiceOrderPhoto.id)).filter(ServiceOrderPhoto.service_order_id==order_id).scalar() or 0
    if existing+len(files)>MAX_PHOTOS_PER_ORDER:
        raise HTTPException(400,f'Limite de {MAX_PHOTOS_PER_ORDER} fotos por ordem de serviço: já existem {existing} e você enviou {len(files)}. Nenhuma foto foi anexada.')
    prepared=[]
    for upload in files:  # valida TUDO antes de gravar qualquer coisa
        name=_clean_filename(upload.filename)
        data=upload.file.read(MAX_PHOTO_BYTES+1)
        if not data: raise HTTPException(400,f'A foto "{name}" está vazia. Nenhuma foto foi anexada.')
        if len(data)>MAX_PHOTO_BYTES: raise HTTPException(400,f'A foto "{name}" passa de {MAX_PHOTO_BYTES//1024//1024} MB. Nenhuma foto foi anexada.')
        kind=_sniff_image(data[:16])
        if not kind: raise HTTPException(400,f'"{name}" não é uma imagem JPEG, PNG ou WebP válida. Nenhuma foto foi anexada.')
        prepared.append((name,kind,data))
    folder=_photo_dir(order_id); written=[]; rows=[]
    try:
        folder.mkdir(parents=True,exist_ok=True)
        for name,kind,data in prepared:
            stored=f'{uuid.uuid4().hex}{PHOTO_EXTENSIONS[kind]}'
            (folder/stored).write_bytes(data); written.append(folder/stored)
            row=ServiceOrderPhoto(service_order_id=order_id,original_name=name,stored_name=stored,content_type=kind,size=len(data)); db.add(row); rows.append(row)
        db.commit()
    except Exception:
        db.rollback(); _remove_photo_files(order_id,written)
        api_logger.exception('Falha ao gravar fotos da OS %s',order_id)
        raise HTTPException(500,'Não foi possível salvar as fotos. Nenhuma foto foi anexada.')
    return {'max':MAX_PHOTOS_PER_ORDER,'items':[_photo_out(order_id,row) for row in rows]}

@app.get('/api/service-orders/{order_id}/photos/{photo_id}/file')
def get_order_photo(order_id:int,photo_id:int,db:Session=Depends(get_db)):
    row=db.get(ServiceOrderPhoto,photo_id)
    if not row or row.service_order_id!=order_id: raise HTTPException(404,'Foto não encontrada')
    if not _STORED_NAME.match(row.stored_name or ''): raise HTTPException(404,'Foto não encontrada')
    path=_photo_dir(order_id)/row.stored_name
    if not path.is_file(): raise HTTPException(404,'O arquivo desta foto não está na pasta de documentos (pode ter sido removido ou não veio no backup).')
    return FileResponse(path,media_type=row.content_type,headers={'X-Content-Type-Options':'nosniff','Cache-Control':'private, max-age=3600','Content-Disposition':'inline'})

@app.delete('/api/service-orders/{order_id}/photos/{photo_id}')
def delete_order_photo(order_id:int,photo_id:int,db:Session=Depends(get_db)):
    row=db.get(ServiceOrderPhoto,photo_id)
    if not row or row.service_order_id!=order_id: raise HTTPException(404,'Foto não encontrada')
    stored=row.stored_name
    db.delete(row); db.commit()
    if _STORED_NAME.match(stored or ''): _remove_photo_files(order_id,[_photo_dir(order_id)/stored])
    return {'deleted':True,'id':photo_id}

@app.post('/api/products/{product_id}/movement')
def movement(product_id:int,kind:str=Query(pattern='^(entrada|saida)$'),quantity:int=Query(gt=0),reason:str=Query(min_length=2),reference:str|None=None,idempotency_key:str|None=Header(default=None,alias='Idempotency-Key'),db:Session=Depends(get_db)):
    source_key=f'manual:{idempotency_key}' if idempotency_key else None
    if source_key:
        old=db.query(Movement).filter_by(source_key=source_key).first()
        if old:
            if (old.product_id,old.kind,old.quantity,old.reason,old.reference)!=(product_id,kind,quantity,reason,reference): raise HTTPException(409,'Idempotency-Key já foi usado com dados diferentes.')
            product=db.get(Product,product_id)
            return {'quantity':product.quantity if product else None,'duplicate_request':True}
    product=db.query(Product).filter_by(id=product_id).with_for_update().first()
    if not product: raise HTTPException(404,'Produto não encontrado')
    new_qty=product.quantity+(quantity if kind=='entrada' else -quantity)
    if new_qty<0: raise HTTPException(409,'Saída bloqueada: estoque insuficiente')
    product.quantity=new_qty; db.add(Movement(product_id=product_id,kind=kind,quantity=quantity,reason=reason,reference=reference,source_key=source_key)); db.commit()
    return {'quantity':product.quantity}

@app.get('/api/movements')
def movements(db:Session=Depends(get_db)): return db.query(Movement).order_by(Movement.id.desc()).limit(1000).all()

IMPORT_ALIASES={
 'name':('produto','nome','descricao','descricaodoproduto'),
 'sku':('codigo','codigosku','sku','referencia'),
 'category':('categoria','grupo'),
 'supplier':('fornecedor','marca'),
 'quantity':('qtdestoque','quantidade','estoque','quantidadeestoque'),
 'cost':('custounitr','custounitario','custo','precocusto'),
 'price':('precovendaunitr','precovendaunitario','precovenda','preco'),
 'min_quantity':('estoqueminimo','qtdminima','quantidademinima'),
}
def normalize_header(value):
    value=unicodedata.normalize('NFKD',str(value or '')).encode('ascii','ignore').decode().lower()
    return re.sub(r'[^a-z0-9]','',value)
def locate_product_header(workbook, preferred_sheet=None):
    sheets=[workbook[preferred_sheet]] if preferred_sheet in workbook.sheetnames else list(workbook.worksheets)
    for ws in sheets:
        for row_idx,row in enumerate(ws.iter_rows(min_row=1,max_row=min(ws.max_row or 1,40),values_only=True),1):
            normalized=[normalize_header(x) for x in row]
            mapping={}
            for field,aliases in IMPORT_ALIASES.items():
                match=next((i for i,h in enumerate(normalized) if h and any(a in h or h in a for a in aliases)),None)
                if match is not None: mapping[field]=str(row[match]).strip()
            if 'name' in mapping and ('quantity' in mapping or 'sku' in mapping):
                return ws,row_idx,[str(x).strip() if x is not None else '' for x in row],mapping
    ws=sheets[0]; rows=ws.iter_rows(min_row=1,max_row=1,values_only=True); row=next(rows,())
    return ws,1,[str(x).strip() if x is not None else '' for x in row],{}

@app.post('/api/import/preview')
async def import_preview(file:UploadFile=File(...),sheet_name:str|None=None):
    if not file.filename.lower().endswith('.xlsx'): raise HTTPException(400,'Envie um arquivo .xlsx')
    content=await file.read()
    try: wb=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    except Exception: raise HTTPException(400,'Arquivo Excel inválido ou corrompido')
    ws,header_row,headers,suggested=locate_product_header(wb,sheet_name)
    rows=list(ws.iter_rows(min_row=header_row+1,values_only=True))
    nonempty=[row for row in rows if any(v is not None for v in row)]
    return {'filename':file.filename,'sheet':ws.title,'sheets':wb.sheetnames,'header_row':header_row,'headers':headers,'suggested_mapping':suggested,'sample':[[str(v) if v is not None else '' for v in row] for row in nonempty[:5]],'rows':len(nonempty)}

@app.post('/api/import/products')
async def import_products(file:UploadFile=File(...),mapping:str|None=None,key:str|None='sku',sheet_name:str|None=None,header_row:int|None=None,db:Session=Depends(get_db)):
    if not file.filename.lower().endswith('.xlsx'): raise HTTPException(400,'Envie um arquivo .xlsx')
    try:
        wb=load_workbook(io.BytesIO(await file.read()),read_only=True,data_only=True)
        ws,detected_row,headers,suggested=locate_product_header(wb,sheet_name)
        if header_row is not None:
            header_row=max(1,header_row); row_values=next(ws.iter_rows(min_row=header_row,max_row=header_row,values_only=True),()); headers=[str(x).strip() if x is not None else '' for x in row_values]; data_start=header_row+1
        else: data_start=detected_row+1
        map_=json.loads(mapping) if mapping else suggested
        indexes={k:headers.index(v) for k,v in map_.items() if k in IMPORT_ALIASES and v in headers}
        if 'name' not in indexes: raise HTTPException(400,'Mapeie a coluna de produto/nome antes de importar')
        created=0; errors=[]; duplicates=[]
        for n,row in enumerate(ws.iter_rows(min_row=data_start,values_only=True),data_start):
            if not any(v is not None for v in row): continue
            data={k:(row[i] if i<len(row) else None) for k,i in indexes.items()}; name=str(data.get('name') or '').strip(); sku=str(data.get('sku') or '').strip() or None
            try:
                if not name: raise ValueError('produto/nome obrigatório')
                if key=='sku' and sku and db.query(Product).filter(Product.sku==sku).first(): duplicates.append(n); continue
                if not sku and db.query(Product).filter(Product.name.ilike(name),Product.category==data.get('category')).first(): duplicates.append(n); continue
                def numeric(field,default=0,integer=False):
                    val=data.get(field)
                    if val is None or val=='': return default
                    if isinstance(val,str): val=float(val.replace('R$','').replace(' ','').replace('.','').replace(',','.'))
                    number=float(val)
                    if number<0: raise ValueError(f'{field} não pode ser negativo')
                    return int(number) if integer else number
                p=Product(name=name,sku=sku,category=str(data.get('category') or '').strip() or None,supplier=str(data.get('supplier') or '').strip() or None,quantity=numeric('quantity',0,True),min_quantity=numeric('min_quantity',0,True),cost=numeric('cost'),price=numeric('price'))
                with db.begin_nested(): db.add(p); db.flush()
                created+=1
            except Exception as e: errors.append({'row':n,'error':str(e)})
        db.commit(); report={'created':created,'duplicates':duplicates,'errors':errors}; db.add(ImportLog(filename=file.filename,result=json.dumps(report,ensure_ascii=False))); db.commit(); return report
    except HTTPException: raise
    except Exception as e: db.rollback(); raise HTTPException(400,f'Falha na importação: {e}')

@app.get('/api/export/products')
def export_products(db:Session=Depends(get_db)):
    wb=Workbook(); ws=wb.active; ws.title='Estoque'; ws.append(['Código','Nome','Categoria','Fornecedor','Quantidade','Mínimo','Custo','Preço'])
    for p in db.query(Product).order_by(Product.name): ws.append([p.sku,p.name,p.category,p.supplier,p.quantity,p.min_quantity,p.cost,p.price])
    buf=io.BytesIO(); wb.save(buf); buf.seek(0); return StreamingResponse(buf,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename=estoque_lmf.xlsx'})

@app.get('/api/reports/summary')
def reports_summary(start:date|None=None,end:date|None=None,status:str|None=None,customer_id:int|None=None,product_id:int|None=None,supplier_id:int|None=None,db:Session=Depends(get_db)):
    sales=db.query(Sale)
    if start: sales=sales.filter(func.date(Sale.sold_at)>=start.isoformat())
    if end: sales=sales.filter(func.date(Sale.sold_at)<=end.isoformat())
    if status: sales=sales.filter(Sale.status==status)
    if customer_id: sales=sales.filter(Sale.customer_id==customer_id)
    if product_id: sales=sales.filter(Sale.id.in_(db.query(SaleItem.sale_id).filter(SaleItem.product_id==product_id)))
    if supplier_id: sales=sales.filter(Sale.id.in_(db.query(SaleItem.sale_id).join(Product,Product.id==SaleItem.product_id).filter(Product.supplier_id==supplier_id)))
    chosen=sales.all(); ids=[s.id for s in chosen]
    top=[]
    if ids:
        top=(db.query(Product.name,func.sum(SaleItem.quantity).label('units'),func.sum(SaleItem.quantity*SaleItem.unit_price-SaleItem.discount).label('gross'))
             .join(SaleItem,SaleItem.product_id==Product.id).filter(SaleItem.sale_id.in_(ids)).group_by(Product.id).order_by(func.sum(SaleItem.quantity).desc()).limit(20).all())
    entries=db.query(FinancialEntry)
    if start: entries=entries.filter(FinancialEntry.competence_on>=start)
    if end: entries=entries.filter(FinancialEntry.competence_on<=end)
    if status: entries=entries.filter(FinancialEntry.status==status)
    if customer_id: entries=entries.filter(FinancialEntry.customer_id==customer_id)
    if product_id: entries=entries.filter(FinancialEntry.sale_id.in_(db.query(SaleItem.sale_id).filter(SaleItem.product_id==product_id)))
    if supplier_id: entries=entries.filter(FinancialEntry.order_id.in_(db.query(Order.id).filter(Order.supplier_id==supplier_id)))
    fin=entries.all()
    supplier_rows=(db.query(Supplier.name,func.count(Order.id).label('orders'),func.sum(Order.total).label('ordered'),func.sum(Order.paid).label('paid'))
        .join(Order,Order.supplier_id==Supplier.id).filter(Order.kind.in_(['compra','fornecedor']))
        .group_by(Supplier.id).order_by(func.sum(Order.total).desc()).limit(20).all())
    return {'filters':{'start':start,'end':end,'status':status},'sales_count':len(chosen),'sales_total':sum(s.total for s in chosen),'estimated_cost':sum(s.estimated_cost for s in chosen),'received':sum(s.received_amount for s in chosen),'pending':sum(max(0,s.total-s.received_amount) for s in chosen),'top_products':[{'product':r.name,'units':r.units,'gross':r.gross} for r in top],
      'low_stock':[{'name':p.name,'quantity':p.quantity,'minimum':p.min_quantity} for p in db.query(Product).filter(Product.active==True,Product.quantity<=Product.min_quantity).order_by(Product.quantity).all()],
      'service_status':{status_:count for status_,count in db.query(ServiceOrder.status,func.count(ServiceOrder.id)).group_by(ServiceOrder.status).all()},
      'active_guarantees':db.query(Guarantee).filter(Guarantee.status=='ativa',Guarantee.ends_on>=date.today()).count(),
      'finance':{'income':sum(e.amount for e in fin if e.kind=='receita'),'expenses':sum(e.amount for e in fin if e.kind=='despesa'),'received':sum(e.amount for e in fin if e.kind=='receita' and e.paid_on),'paid':sum(e.amount for e in fin if e.kind=='despesa' and e.paid_on),'pending':sum(e.amount for e in fin if not e.paid_on)},
      'supplier_orders_pending':db.query(Order).filter(Order.kind=='compra',Order.status.notin_(['Concluído','Cancelado'])).count(),
      'supplier_performance':[{'supplier':r.name,'orders':r.orders,'ordered':r.ordered or 0,'paid':r.paid or 0,'outstanding':(r.ordered or 0)-(r.paid or 0)} for r in supplier_rows]}

@app.get('/api/reports/export')
def export_report(start:date|None=None,end:date|None=None,status:str|None=None,customer_id:int|None=None,product_id:int|None=None,supplier_id:int|None=None,db:Session=Depends(get_db)):
    query=db.query(Sale)
    if start: query=query.filter(func.date(Sale.sold_at)>=start.isoformat())
    if end: query=query.filter(func.date(Sale.sold_at)<=end.isoformat())
    if status: query=query.filter(Sale.status==status)
    if customer_id: query=query.filter(Sale.customer_id==customer_id)
    if product_id: query=query.filter(Sale.id.in_(db.query(SaleItem.sale_id).filter(SaleItem.product_id==product_id)))
    if supplier_id: query=query.filter(Sale.id.in_(db.query(SaleItem.sale_id).join(Product,Product.id==SaleItem.product_id).filter(Product.supplier_id==supplier_id)))
    wb=Workbook(); ws=wb.active; ws.title='Vendas'; ws.append(['Venda','Data','Status','Total','Custo estimado','Recebido','Pendente'])
    for sale in query.order_by(Sale.sold_at): ws.append([sale.number,sale.sold_at.strftime('%Y-%m-%d %H:%M'),sale.status,sale.total,sale.estimated_cost,sale.received_amount,max(0,sale.total-sale.received_amount)])
    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename=relatorio_lmf.xlsx'})

def make_backup():
    return backup_service.create_backup()
@app.post('/api/backups')
def backup():
    try: return {'file':make_backup().name,'status':'sucesso'}
    except Exception as e: raise HTTPException(500,f'Falha ao criar backup: {e}')
@app.get('/api/backups')
def backups(): return {'items':backup_service.list_backups(),'importable':backup_service.importable_backups(),'paths':{'data':str(DATA_DIR),'database':str(DB_PATH),'backups':str(BACKUP_DIR),'import':str(IMPORT_BACKUP_DIR)},'backup_path_notice':_backup_path_notice,'settings':backup_service.read_settings(),'last_status':backup_service.last_status()}
@app.post('/api/backups/{filename}/restore')
def restore(filename:str,request:Request,confirm:bool=Query(False),allow_incomplete:bool=Query(False)):
    if not confirm: raise HTTPException(400,'Confirme a restauração com confirm=true')
    if Path(filename).name!=filename: raise HTTPException(404,'Backup não encontrado')
    authorization.require_authorization_standalone(request,'backup.restaurar',filename)
    src=BACKUP_DIR/filename
    try: return backup_service.restore_backup(src,backup_service.schema_revision(),allow_incomplete)
    except Exception as e: raise HTTPException(400,f'Falha na restauração: {e}')

@app.post('/api/backups/import/{filename}')
def import_backup(filename:str):
    try:
        path=backup_service.copy_imported_backup(filename)
        return {'name':path.name,'status':'validado e importado'}
    except Exception as e: raise HTTPException(400,f'Não foi possível importar: {e}')

@app.delete('/api/backups/{filename}')
def remove_backup(filename:str,confirm:bool=Query(False)):
    if not confirm: raise HTTPException(400,'Confirme a exclusão com confirm=true')
    try: backup_service.delete_backup(filename); return {'status':'excluído'}
    except Exception as e: raise HTTPException(400,str(e))

@app.put('/api/backups/settings')
def save_backup_settings(payload:dict):
    try:
        backup_time=payload.get('backup_time','02:00'); datetime.strptime(backup_time,'%H:%M')
        count=max(1,min(1000,int(payload.get('retention_count',30))))
        days=max(1,min(3650,int(payload.get('retention_days',180))))
        return backup_service.write_settings({'backup_time':backup_time,'retention_count':count,'retention_days':days})
    except Exception as e: raise HTTPException(400,f'Configuração inválida: {e}')

@app.post('/api/backups/open-folder/{folder}')
def open_backup_folder(folder:str):
    target={'backups':BACKUP_DIR,'import':IMPORT_BACKUP_DIR}.get(folder)
    if target is None: raise HTTPException(404,'Pasta não encontrada')
    try:
        import os
        target.mkdir(parents=True,exist_ok=True)
        if os.name=='nt': os.startfile(str(target.resolve()))
        else: raise OSError('Abertura automática disponível somente no Windows.')
        return {'status':'opened','path':str(target.resolve())}
    except OSError as e: raise HTTPException(500,f'Não foi possível abrir a pasta {target}: {e}')
